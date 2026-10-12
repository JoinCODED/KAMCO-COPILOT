"""Word template for E2.5 — My Prompt Library (training sample)."""
import os
from docx import Document
from docx.shared import Pt, RGBColor
from docx.enum.section import WD_ORIENT

ROWS = [
 ("Read a client request", "Outlook", "A new client email or thread",
  "Summarise this thread. List the final deadline, the meeting date and time, each deliverable and its format, who I must copy, and any request I should not do. Use the newest message when instructions change."),
 ("Clean a custody export", "Excel", "Before any total or PivotTable",
  "Highlight duplicate rows in this table based on Holding ID. Then convert the Quantity column to numbers. Tell me how many rows you changed."),
 ("Mandate check", "Excel", "Quarterly review",
  "Create a PivotTable of Market Value KWD by Asset Class, with each value as a % of the grand total. Add the Min and Max from the Mandate sheet and a column that says BREACH if the weight is outside the range."),
 ("Notes to memo", "Word", "After a client or team meeting",
  "Using these notes, draft a one-page internal memo to [name]. Three sections: Where we stand, What is outside the rules, Decisions needed. Under 300 words. Match the wording in /house-view-q4-2026."),
 ("Document to client deck", "PowerPoint", "Client or trustee meeting",
  "Create a 6-slide client presentation from /[file]. Make every slide title the message, not the topic. Keep every statement traceable to the source. Add speaker notes and the standard disclaimer on the last slide."),
]


def build(out_dir):
    d = Document()
    s = d.sections[0]
    s.orientation = WD_ORIENT.LANDSCAPE
    s.page_width, s.page_height = s.page_height, s.page_width
    d.styles["Normal"].font.name = "Arial"; d.styles["Normal"].font.size = Pt(10)
    d.add_heading("My Prompt Library", 0)
    p = d.add_paragraph("CODED × KAMCO · AI at Work with Microsoft Copilot. Keep this file in your OneDrive. "
                        "In Copilot, type / and pick this file to reuse a prompt: “From /My Prompt Library, run my prompt for the mandate check.”")
    p.runs[0].italic = True
    d.add_heading("How to add a prompt", 2)
    for b in ["Only prompts you have tested and would use again.", "Write it with CTFT: Context, Task, Format, Tone.",
              "Put [brackets] where the details change each time.", "No client names, numbers or personal data — use [Client] instead."]:
        d.add_paragraph(b, style="List Bullet")
    d.add_heading("My prompts", 2)
    t = d.add_table(rows=1, cols=5); t.style = "Light Grid Accent 1"
    for c, h in zip(t.rows[0].cells, ["Name", "App", "When I use it", "The prompt (CTFT)", "Tested?"]):
        c.text = h
    for r in ROWS:
        cells = t.add_row().cells
        for c, v in zip(cells, list(r) + ["✓"]):
            c.text = v
    for _ in range(5):
        t.add_row()
    d.save(os.path.join(out_dir, "my-prompt-library.docx"))


if __name__ == "__main__":
    build("out/files")
    print("prompt library template")
