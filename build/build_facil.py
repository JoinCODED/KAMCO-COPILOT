import json, os, shutil, html as H
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import FormulaRule
from build_pages import panel_page, res
from emails import MURJAN_THREAD, PROJECT_A, PROJECT_B, thread_text
from scenarios import MANDATE

A = json.load(open("out/answers.json"))
OUTF = "out/facilitator"


def kwd(x):
    return "KWD {:,.0f}".format(x)


def tbl(rows, head=None):
    st = 'style="width:100%;border-collapse:collapse;font-size:14px;margin-top:10px"'
    td = 'style="padding:8px 10px;border-bottom:1px solid var(--line);color:var(--ink-dim);vertical-align:top"'
    th = 'style="padding:8px 10px;border-bottom:1px solid var(--line-2);color:var(--ink);text-align:left;font-family:var(--mono);font-size:11px;letter-spacing:.08em;text-transform:uppercase"'
    h = "<tr>" + "".join(f"<th {th}>{x}</th>" for x in head) + "</tr>" if head else ""
    return f"<table {st}>{h}" + "".join("<tr>" + "".join(f"<td {td}>{c}</td>" for c in r) + "</tr>" for r in rows) + "</table>"


def mailbox(t):
    return '<div class="mailtxt">' + H.escape(t) + "</div>"


def run_of_show():
    d1 = [("09:00", "MAP test (pre)", "Aya", "Ali: sign-in checks, backup accounts"),
          ("09:30", "Opening · Copilot, honestly · Lab 1", "Ali", "Aya: Work vs Web demo live"),
          ("10:30", "Prompting for investment work · Lab 2", "Aya", "Ali: walk the room"),
          ("11:30", "Read the request in Outlook · Labs 3–4", "Aya", "Ali: watch for the confidentiality trap"),
          ("13:00", "Excel: clean the data · Lab 5", "Aya", "Ali: support slow laptops"),
          ("14:00", "Excel: find the story · Labs 6–7", "Aya", "Ali: support"),
          ("15:00", "Hands-on (observed)", "Both", "Write Day 1 observation notes in the scoring sheet")]
    d2 = [("09:00", "Memo or deck · Word · Labs 1–2", "Ali", "Aya: support"),
          ("10:00", "PowerPoint · Lab 3", "Ali", "Aya: support"),
          ("11:00", "Reply & the four checks · Labs 4–5", "Aya", "Ali: support"),
          ("13:00", "Agent Builder demo · Lab 6 · project briefing", "Aya", "Ali: send project emails at 13:30"),
          ("13:30", "Project (graded)", "Both", "Answer questions about the tools only — not the analysis"),
          ("15:00", "MAP test (post)", "Ali", "Aya: start scoring"),
          ("15:30", "Certificates & close", "Both", "Visitors (Hashim, Ahmad, BD) welcome here")]
    before = ["<b>Licences:</b> all 35 trainees have <b>paid Microsoft 365 Copilot</b> (not only free Copilot Chat) with Outlook, Excel, Word and PowerPoint active.",
              "<b>Agent Builder</b> switched on in KAMCO's tenant (ask Ahmed's IT team). If not, the agent segment stays a demo.",
              "<b>Backup accounts:</b> ~5 CODED accounts with Copilot, tested, credentials ready for the onboarding email. The 10 backup laptops signed in.",
              "<b>Trainee names and KAMCO emails</b> at least 3 days before (needed to send the scenario emails).",
              "<b>IT allows email</b> from the CODED evaluation mailbox into KAMCO inboxes (send a test).",
              "<b>Evaluation mailbox</b> created; the scenario and project emails are sent <b>from</b> it, so replies land there.",
              "<b>MAP test links</b> (pre and post) ready.",
              "<b>Certificates</b> printed with the final attendee names.",
              "<b>Rehearsal with Ali:</b> run both decks once; score one sample project together and compare."]
    morning = ["<b>Day 1, 08:45</b> — send the Al-Murjan thread (below) to every trainee, with <b>amfo-q3-custody-export.xlsx</b> attached.",
               "<b>Day 1, 09:00</b> — open the deck on the room screen, MAP test link on the board.",
               "<b>Day 2, 13:30</b> — send Project A to odd seats and Project B to even seats (texts on the project page). Note each trainee's version in the scoring sheet."]
    rows1 = tbl([list(r) for r in d1], ["Time", "Block", "Lead", "Support / notes"])
    rows2 = tbl([list(r) for r in d2], ["Time", "Block", "Lead", "Support / notes"])
    return panel_page("Run of show & readiness — KAMCO Copilot (facilitator)", "Facilitator · Aya & Ali", "Run of show & readiness",
        "Who leads each block, what must be true before Sunday, and the email you send on Day 1 morning.",
        [("Day 1 · Sun 11 Oct", rows1), ("Day 2 · Mon 12 Oct", rows2),
         ("This week · readiness checklist", '<ol class="steps2">' + "".join(f"<li>{x}</li>" for x in before) + "</ol>"),
         ("On the day · emails to send", '<ol class="steps2">' + "".join(f"<li>{x}</li>" for x in morning) + "</ol>"),
         ("Day 1 email · Al-Murjan thread (send at 08:45)", "<p>Subject: <b>" + H.escape(MURJAN_THREAD["subject"]) + "</b> · Attach: <b>amfo-q3-custody-export.xlsx</b>. The .eml file in this kit has the attachment built in.</p>" + mailbox(thread_text(MURJAN_THREAD))),
         ("Instructor demo files", '<div class="res-list">' + res("awfo-q3-custody-export.xlsx", "Al-Waha custody export (instructor demos)", "Excel") + res("00-awfo-instructor-demo-request.eml", "Al-Waha thread (instructor demos)", "Email") + "</div>")])


def project_page():
    def ref(k):
        a = A[k]
        al = ", ".join(f"{x} {v}%" for x, v in a["alloc_pct"].items())
        return tbl([["Raw rows / duplicates / holdings", f"{a['raw_rows']} / {a['duplicates']} / {a['holdings']}"],
                    ["Missing prices", ", ".join(a["missing_tickers"])],
                    ["Value at 30 Sep", kwd(a["total_sep"])], ["Value if duplicates kept (common error)", kwd(a["wrong_if_dups_kept"])],
                    ["Q3 change", f"{kwd(a['q3_change'])} · {a['q3_return_pct']}%"], ["Allocation", al],
                    ["Mandate breach", ", ".join(a["breaches"]) or "None"],
                    ["Issuer check", "; ".join(f"{n} {v}%" for n, v in a["issuer_top"]) + (" → <b>breach: " + ", ".join(n for n, v in a["issuer_breaches"]) + "</b>" if a["issuer_breaches"] else " → <b>no breach</b>")],
                    ["Top 3 contributors", ", ".join(f"{n} (+{v:,})" for n, v in a["top3"])]])
    rub = tbl([["Workflow complete", "30", "Read the email, analysed the data, produced the deliverable, replied with it attached by 2:50 pm."],
               ["Accurate numbers", "25", "Value, Q3 change, allocation, breaches and issuer check match the reference (±0.1 pt)."],
               ["Deliverable quality", "20", "A: one page, answer first. B: 5 slides, message titles, disclaimer included."],
               ["AI review & judgement", "15", "“What I corrected” line shows at least one real correction or check."],
               ["Prompt quality", "10", "Ask to see 2–3 prompts in Copilot history: specific, sourced, CTFT."]], ["Criterion", "Points", "Full marks when…"]) + "<p><b>Finished early?</b> Trainees may build the other format too (A → a 3-slide deck, B → a one-page memo). It is not graded.</p>"
    return panel_page("Project & scoring — KAMCO Copilot (facilitator)", "Facilitator · Day 2 project", "Project & scoring",
        "Send at 13:30 from the evaluation mailbox: version A to odd seats, B to even seats. Replies arrive with a timestamp. Score with the rubric in the scoring sheet.",
        [("Rubric (out of 100)", rub + "<p>Scores are ranked in the scoring sheet. Day 1 observation notes and the MAP gain settle borderline cases and ties. Score one sample together first so you both mark the same way.</p>"),
         ("Version A · " + PROJECT_A["client"] + " · Word memo", f"<p>Subject: <b>{H.escape(PROJECT_A['subject'])}</b> · From: {PROJECT_A['frm']} · Attach: <b>{PROJECT_A['attach']}</b></p>" + mailbox(PROJECT_A["body"]) + "<p><b>Reference answers</b></p>" + ref("bayan")),
         ("Version B · " + PROJECT_B["client"] + " · PowerPoint deck", f"<p>Subject: <b>{H.escape(PROJECT_B['subject'])}</b> · From: {PROJECT_B['frm']} · Attach: <b>{PROJECT_B['attach']}</b></p>" + mailbox(PROJECT_B["body"]) + "<p><b>Reference answers</b></p>" + ref("sanad")),
         ("Files", '<div class="res-list">' + res("project-A-al-bayan-request.eml", "Project A email (.eml, attachment built in)", "Email") + res("abh-q3-custody-export.xlsx", "Project A custody export", "Excel")
          + res("project-B-dar-al-sanad-request.eml", "Project B email (.eml, attachment built in)", "Email") + res("dast-q3-custody-export.xlsx", "Project B custody export", "Excel") + res("kamco-p2-scoring-sheet.xlsx", "Scoring sheet (35 trainees)", "Excel") + "</div>")])


def scoring_sheet(path):
    wb = Workbook(); ws = wb.active; ws.title = "Scores"
    F = lambda **k: Font(name="Arial", **k)
    blue = F(color="0000FF"); bold = F(bold=True); hdrfill = PatternFill("solid", fgColor="D9D9D9"); yel = PatternFill("solid", fgColor="FFF2CC")
    ws["A1"] = "KAMCO Invest · AI at Work with Microsoft Copilot · Scoring sheet"; ws["A1"].font = F(bold=True, size=13)
    ws["A2"] = "Blue cells = enter values. Black cells = formulas. Seats for Program 3 →"; ws["A2"].font = F(italic=True, color="666666")
    ws["J2"] = 30; ws["J2"].font = blue; ws["J2"].fill = yel
    ws["K2"] = "Pass mark →"; ws["K2"].font = F(italic=True, color="666666"); ws["L2"] = 60; ws["L2"].font = blue; ws["L2"].fill = yel
    hdr = ["#", "Trainee name", "Department", "Version (A/B)", "MAP pre (%)", "MAP post (%)", "MAP gain (pts)",
           "Workflow (30)", "Accuracy (25)", "Deliverable (20)", "AI review (15)", "Prompts (10)", "Project total (100)",
           "Rank", "Recommendation", "Day 1 observation notes", "Reply time"]
    for j, h in enumerate(hdr, 1):
        c = ws.cell(row=4, column=j, value=h); c.font = bold; c.fill = hdrfill; c.alignment = Alignment(wrap_text=True, vertical="center")
    caps = {8: 30, 9: 25, 10: 20, 11: 15, 12: 10}
    for i in range(35):
        r = 5 + i
        ws.cell(row=r, column=1, value=i + 1).font = F()
        for j in (2, 3, 4, 5, 6, 8, 9, 10, 11, 12, 16, 17):
            ws.cell(row=r, column=j).font = blue
        ws.cell(row=r, column=7, value=f'=IF(AND(ISNUMBER(E{r}),ISNUMBER(F{r})),F{r}-E{r},"")').font = F()
        ws.cell(row=r, column=13, value=f'=IF(COUNT(H{r}:L{r})=0,"",SUM(H{r}:L{r}))').font = F(bold=True)
        # rank on project total, MAP post breaks ties
        ws.cell(row=r, column=14, value=f'=IF(M{r}="","",COUNTIF($M$5:$M$39,">"&M{r})+COUNTIFS($M$5:$M$39,M{r},$F$5:$F$39,">"&N(F{r}))+1)').font = F()
        ws.cell(row=r, column=15, value=f'=IF(M{r}="","",IF(AND(M{r}>=$L$2,N{r}<=$J$2),"Recommend",IF(M{r}>=$L$2,"Passed - over seat limit","Review / retake")))').font = F()
        for col, cap in caps.items():
            pass
    # example row
    ex = ["Example Trainee (delete)", "Wealth Management", "A", 55, 80, None, 26, 20, 16, 12, 8]
    for j, v in zip([2, 3, 4, 5, 6, 7, 8, 9, 10, 11, 12], ex):
        if v is not None:
            ws.cell(row=5, column=j, value=v)
    ws.cell(row=5, column=16, value="Cleaned before counting; caught the other-family request.")
    dv = DataValidation(type="list", formula1='"A,B"', allow_blank=True); ws.add_data_validation(dv); dv.add("D5:D39")
    for col, cap in caps.items():
        L = get_column_letter(col)
        d = DataValidation(type="decimal", operator="between", formula1="0", formula2=str(cap), allow_blank=True, error="Score must be between 0 and %d" % cap)
        ws.add_data_validation(d); d.add(f"{L}5:{L}39")
    ws.conditional_formatting.add("O5:O39", FormulaRule(formula=['$O5="Recommend"'], fill=PatternFill("solid", fgColor="C6EFCE")))
    ws.conditional_formatting.add("O5:O39", FormulaRule(formula=['$O5="Review / retake"'], fill=PatternFill("solid", fgColor="FFC7CE")))
    widths = [4, 26, 18, 10, 10, 10, 10, 11, 11, 11, 11, 11, 12, 7, 24, 44, 11]
    for j, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.row_dimensions[4].height = 32; ws.freeze_panes = "C5"
    r = 41
    ws.cell(row=r, column=2, value="Summary").font = bold
    for k, (lab, f) in enumerate([("Trainees scored", '=COUNT(M5:M39)'), ("Average project score", '=IF(COUNT(M5:M39)=0,"",AVERAGE(M5:M39))'),
                                  ("Average MAP gain (pts)", '=IF(COUNT(G5:G39)=0,"",AVERAGE(G5:G39))'), ("Recommended for Program 3", '=COUNTIF(O5:O39,"Recommend")')]):
        ws.cell(row=r + 1 + k, column=2, value=lab).font = F(); c = ws.cell(row=r + 1 + k, column=4, value=f); c.font = F(bold=True)
        if "Average" in lab: c.number_format = "0.0"
    # rubric + reference sheets
    rb = wb.create_sheet("Rubric")
    rows = [("Criterion", "Points", "Full marks", "Half marks", "Zero"),
            ("Workflow complete", 30, "All four steps; reply with file attached by 2:50 pm", "A step missing or reply late", "No reply"),
            ("Accurate numbers", 25, "Value, Q3 change, allocation, breaches and issuer check match the reference", "One or two errors", "Duplicates kept / major errors"),
            ("Deliverable quality", 20, "A: one page, answer first. B: 5 slides, message titles, disclaimer", "Right format, weak structure", "Wrong format"),
            ("AI review & judgement", 15, "'What I corrected' shows a real correction", "Generic note", "No review shown"),
            ("Prompt quality", 10, "Specific, sourced, CTFT prompts", "Some structure", "Generic prompts")]
    for i, row in enumerate(rows, 1):
        for j, v in enumerate(row, 1):
            c = rb.cell(row=i, column=j, value=v); c.font = bold if i == 1 else F(); c.alignment = Alignment(wrap_text=True, vertical="top")
            if i == 1: c.fill = hdrfill
    for j, w in enumerate([22, 8, 50, 30, 28], 1):
        rb.column_dimensions[get_column_letter(j)].width = w
    rf = wb.create_sheet("Reference answers")
    hdr = ["Measure", "Version A · " + A["bayan"]["client"], "Version B · " + A["sanad"]["client"]]
    for j, h in enumerate(hdr, 1):
        c = rf.cell(row=1, column=j, value=h); c.font = bold; c.fill = hdrfill
    def r2(k):
        a = A[k]
        return [a["raw_rows"], a["duplicates"], a["holdings"], ", ".join(a["missing_tickers"]), a["total_sep"], a["wrong_if_dups_kept"], a["q3_change"], a["q3_return_pct"] / 100,
                a["alloc_pct"]["Equity"] / 100, a["alloc_pct"]["Fixed Income"] / 100, a["alloc_pct"]["Real Estate"] / 100, a["alloc_pct"]["Cash"] / 100,
                ", ".join(a["breaches"]) or "None", "; ".join(f"{n} {v}%" for n, v in a["issuer_top"]), ", ".join(n for n, v in a["issuer_breaches"]) or "None",
                ", ".join(n for n, v in a["top3"])]
    labs = ["Raw rows", "Duplicates", "Holdings after cleaning", "Missing prices", "Value 30 Sep (KWD)", "Value if duplicates kept (KWD) — common error", "Q3 change (KWD)", "Q3 change (%)",
            "Equity weight", "Fixed income weight", "Real estate weight", "Cash weight", "Mandate breach", "Top 3 issuers", "Issuer breach (>10%)", "Top 3 contributors"]
    va, vb = r2("bayan"), r2("sanad")
    for i, (l, x, y) in enumerate(zip(labs, va, vb), 2):
        rf.cell(row=i, column=1, value=l).font = F()
        for j, v in ((2, x), (3, y)):
            c = rf.cell(row=i, column=j, value=v); c.font = F()
            if isinstance(v, float) and v < 1: c.number_format = "0.0%"
            elif isinstance(v, (int, float)) and v > 1000: c.number_format = "#,##0"
    rf.cell(row=len(labs) + 3, column=1, value="Source: generated training data (fictional). USD→KWD 0.3065; mandate ranges on each file's Mandate sheet.").font = F(italic=True, color="666666")
    for j, w in enumerate([42, 46, 46], 1):
        rf.column_dimensions[get_column_letter(j)].width = w
    for wsx in wb.worksheets:
        for row in wsx.iter_rows():
            for c in row:
                if c.font and c.font.name != "Arial":
                    c.font = Font(name="Arial", bold=c.font.bold, italic=c.font.italic, color=c.font.color, size=c.font.size)
    wb.save(path)


if __name__ == "__main__":
    os.makedirs(OUTF, exist_ok=True)
    open(os.path.join(OUTF, "run-of-show.html"), "w").write(run_of_show())
    open(os.path.join(OUTF, "project-and-scoring.html"), "w").write(project_page())
    scoring_sheet(os.path.join(OUTF, "kamco-p2-scoring-sheet.xlsx"))
    for f in ["awfo-q3-custody-export.xlsx", "00-awfo-instructor-demo-request.eml", "project-A-al-bayan-request.eml", "abh-q3-custody-export.xlsx",
              "project-B-dar-al-sanad-request.eml", "dast-q3-custody-export.xlsx", "amfo-q3-custody-export.xlsx", "01-amfo-q3-review-request.eml", "house-view-q4-2026.docx"]:
        shutil.copy("out/files/" + f, OUTF)
    print(sorted(os.listdir(OUTF)))
