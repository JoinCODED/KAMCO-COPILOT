import os
from email.message import EmailMessage
from docx import Document
from docx.shared import Pt, RGBColor
from emails import MURJAN_THREAD, WAHA_THREAD, PROJECT_A, PROJECT_B, thread_text

OUT = "out/files"
os.makedirs(OUT, exist_ok=True)

DISCLAIMER = ("This document is for information only and does not constitute an offer or a recommendation to buy or sell any security. "
              "Past performance is not a reliable indicator of future results. The value of investments can go down as well as up.")


def eml(name, subject, frm, body, attach_name=None):
    m = EmailMessage()
    m["Subject"] = subject
    m["From"] = frm + " <training@kamco-training.example>"
    m["To"] = "you@kamco-training.example"
    m.set_content(body)
    if attach_name:
        data = open(os.path.join(OUT, attach_name), "rb").read()
        m.add_attachment(data, maintype="application", subtype="vnd.openxmlformats-officedocument.spreadsheetml.sheet", filename=attach_name)
    open(os.path.join(OUT, name), "wb").write(bytes(m))


def house_view():
    d = Document()
    st = d.styles["Normal"]; st.font.name = "Arial"; st.font.size = Pt(10.5)
    d.add_heading("Investment House View - Q4 2026", 0)
    p = d.add_paragraph("TRAINING SAMPLE - fictional content prepared for the AI at Work with Microsoft Copilot program. Not investment advice.")
    p.runs[0].italic = True
    d.add_heading("1. Summary", 1)
    d.add_paragraph("We stay neutral on risk going into Q4. Equity markets had a strong Q3, led by technology and Gulf banks. "
                    "Portfolios that drifted above their equity range should be brought back inside it. We prefer adding high-quality "
                    "fixed income gradually, and we stay cautious on real estate.")
    d.add_heading("2. Asset class views", 1)
    t = d.add_table(rows=1, cols=3); t.style = "Light Grid Accent 1"
    for c, h in zip(t.rows[0].cells, ["Asset class", "View", "What this means for portfolios"]):
        c.text = h
    rows = [("Equity", "Neutral", "Trim positions back to the mandate range after the Q3 rally. Keep quality Gulf banks, but respect issuer limits."),
            ("Fixed Income", "Positive", "Add gradually to GCC sovereign and high-grade issuers. Keep average maturity around 5 years."),
            ("Real Estate", "Cautious", "Do not add. Hold existing positions within the range. REIT prices remain under pressure."),
            ("Cash", "Neutral", "Keep 5-10% for liquidity. Large cash balances should be invested in line with the mandate.")]
    for r in rows:
        cells = t.add_row().cells
        for c, v in zip(cells, r):
            c.text = v
    d.add_heading("3. Concentration", 1)
    d.add_paragraph("No single issuer should exceed 10% of a client portfolio. Equity and bonds from the same issuer count together. "
                    "Where a strong rally pushes an issuer above the limit, we recommend trimming back below 10% at the next rebalance.")
    d.add_heading("4. Rules for client communication", 1)
    for b in ["Every figure must match the approved analysis. Show KWD values without decimals and percentages to one decimal place.",
              "Never promise or imply a guaranteed return.",
              "Never include another client's name, holdings or results.",
              "Copy Compliance on any document that goes to a client.",
              "Client-facing documents must end with the standard disclaimer below. Internal memos do not need it."]:
        d.add_paragraph(b, style="List Bullet")
    d.add_heading("5. Standard disclaimer (client documents only)", 1)
    p = d.add_paragraph(DISCLAIMER); p.runs[0].italic = True
    d.save(os.path.join(OUT, "house-view-q4-2026.docx"))


if __name__ == "__main__":
    house_view()
    eml("01-amfo-q3-review-request.eml", MURJAN_THREAD["subject"], "Faisal Al-Mutairi", thread_text(MURJAN_THREAD), MURJAN_THREAD["attach"])
    eml("00-awfo-instructor-demo-request.eml", WAHA_THREAD["subject"], "Reem Al-Qattan", thread_text(WAHA_THREAD), WAHA_THREAD["attach"])
    eml("project-A-al-bayan-request.eml", PROJECT_A["subject"], PROJECT_A["frm"].split(" (")[0], PROJECT_A["body"], PROJECT_A["attach"])
    eml("project-B-dar-al-sanad-request.eml", PROJECT_B["subject"], PROJECT_B["frm"].split(" (")[0], PROJECT_B["body"], PROJECT_B["attach"])
    print(sorted(os.listdir(OUT)))
