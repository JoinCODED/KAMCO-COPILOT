"""Knowledge files for the field agents (E2.8). Fictional KAMCO training versions."""
import os
from docx import Document
from docx.shared import Pt

OUT = "out/files"
DOCS = {
 "kamco-hr-policies-2026.docx": ("KAMCO HR Policies 2026 (training)", [
   ("1. Annual leave", ["30 working days a year after probation.", "No annual leave during probation, except up to 5 days of emergency leave with your manager's approval.",
                        "You may carry over up to 10 unused days to the next year. Carried-over days must be used by 31 March. (This replaces the 2024 handbook, which allowed 15 days.)"]),
   ("2. Probation", ["Probation is 6 months for all new joiners."]),
   ("3. Sick leave", ["Up to 15 days a year on full pay.", "A medical certificate is needed for any absence longer than 2 days."]),
   ("4. Hybrid work", ["After probation, up to 2 remote days a week with your manager's approval.", "Client-facing staff must be in the office on days with client meetings."]),
   ("5. Questions", ["Anything not covered here: email hr@kamco-training.example. HR never shares one employee's information with another employee."])]),
 "kamco-hr-handbook-2024.docx": ("KAMCO Employee Handbook 2024 (training)", [
   ("Leave carry-over", ["Employees may carry over up to 15 unused leave days to the next year."]),
   ("Working hours", ["Sunday to Thursday, 8:00 to 16:00."])]),
 "kamco-expenses-policy.docx": ("KAMCO Expenses and Travel Policy (training)", [
   ("1. Approvals", ["Up to KWD 250: your line manager.", "KWD 250 to KWD 1,000: your department head.", "Above KWD 1,000: the CFO."]),
   ("2. Receipts", ["A receipt is needed for every item above KWD 10.", "Submit claims within 30 days of the expense. Late claims need the department head's approval."]),
   ("3. Meals and travel", ["Local business meals: up to KWD 15 per person.", "Per diem: GCC KWD 60 a day, other countries KWD 90 a day.",
                            "Business class only for flights over 6 hours, and only for VP grade and above."]),
   ("4. Never reimbursed", ["Alcohol, personal entertainment, traffic fines and family members' costs."])]),
 "kamco-compliance-rules.docx": ("KAMCO Client Communication and Compliance Rules (training)", [
   ("1. Words we never use", ["Never use: guaranteed, risk-free, assured return, cannot lose.", "Past performance must always be followed by: Past performance is not a guide to future returns."]),
   ("2. Client data", ["Never include another client's name, holdings or results.", "Client details stay inside Compliance and the client's own relationship manager."]),
   ("3. Approvals", ["Any document or message that goes to more than one client, or to the public, needs Compliance approval first."]),
   ("4. Gifts", ["Gifts from or to clients above KWD 50 must be registered with Compliance within 5 business days."]),
   ("5. Complaints", ["Acknowledge within 2 business days. Resolve within 15 calendar days."])]),
 "kamco-ops-procedures.docx": ("KAMCO Operations Procedures (training)", [
   ("1. Settlement", ["Standard settlement is T+2 business days. Friday and Saturday are closed.", "Trade instructions must reach Operations by 12:00 on trade date + 1."]),
   ("2. Failed trades", ["Any trade still open more than 5 calendar days after the expected date is escalated to the COO."]),
   ("3. Bank details (SSI) changes", ["Never change bank or settlement details because of an email alone.",
                                      "Every change needs a call-back to a phone number already on file — never a number in the email — and a second approver in Operations."]),
   ("4. Client data", ["Reports outside Operations use trade IDs and counterparty names only — no client names."])]),
 "kgof-fund-rules.docx": ("Gulf Opportunities Fund — Key Rules (training prospectus extract)", [
   ("1. Limits", ["Single stock: max 10% of NAV.", "Single sector: max 35% of NAV.", "Cash: 2% to 10% of NAV."]),
   ("2. Dealing", ["Dealing day: every Thursday.", "Subscription and redemption cut-off: 12:00 on Wednesday.", "Redemptions are paid within 4 business days of the dealing day."]),
   ("3. Fees", ["Management fee: 1.5% a year. No performance fee."]),
   ("4. Communication", ["Never promise or imply future returns.", "Past performance is not a guide to future returns."])]),
 "kamco-research-guidelines.docx": ("KAMCO Research Guidelines (training)", [
   ("1. Valuation", ["Use the peer median EV/EBITDA on the last full year's EBITDA. Count each peer once and exclude peers without a number.",
                     "View: upside above +15% = Attractive; −10% to +15% = Neutral; below −10% = Unattractive."]),
   ("2. Blackout", ["No research is published in the 5 business days before a company reports its results."]),
   ("3. Information barrier", ["Research is never shared with Corporate Finance or any client before it is published."]),
   ("4. Language", ["No price targets in client documents. No promises about returns."])]),
}


def build(out_dir=OUT):
    for fn, (title, secs) in DOCS.items():
        d = Document(); d.styles["Normal"].font.name = "Arial"; d.styles["Normal"].font.size = Pt(10.5)
        d.add_heading(title, 0)
        p = d.add_paragraph("TRAINING SAMPLE — fictional content for the AI at Work with Microsoft Copilot program."); p.runs[0].italic = True
        for h, bullets in secs:
            d.add_heading(h, 1)
            for b in bullets:
                d.add_paragraph(b, style="List Bullet")
        d.save(os.path.join(out_dir, fn))
    return list(DOCS)


if __name__ == "__main__":
    print(build())
